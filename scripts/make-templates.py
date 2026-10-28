"""Builds the two upload formats in public/ and src/lib/listingFormat.ts.

  python3 scripts/make-templates.py

Global Database Format for Realtors: listing-template.xlsx/.csv (+ listing-example.csv)
Seller's Format for Individual Sellers: seller-template.xlsx/.csv
Column names must stay matched by src/lib/uploadParse.ts.
"""
import csv, json
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

LONG_TERM = "Long-term rentals only: 6 months or more. No holiday lets or short stays."

REALTOR = [
 ("listing_type","Optional","sale or rent. Blank means sale. " + LONG_TERM),
 ("price","Required","Asking price, or the monthly rent for rentals. 450000, 450,000, RM 450,000 and 1.2m all work."),
 ("currency","Optional","3-letter code: MYR, SGD, USD, GBP, EUR, AUD, THB… Read from the price or location if left blank."),
 ("location","Required*","Town and country, e.g. Ipoh, Perak, Malaysia. *Or leave blank and fill latitude and longitude."),
 ("latitude","Optional","Map pin, decimal degrees, e.g. 4.5975. Puts the listing on the exact spot on the world map."),
 ("longitude","Optional","Map pin, decimal degrees, e.g. 101.0901."),
 ("contact_name","Required","Seller's, landlord's or realtor's name, shown to buyers and tenants."),
 ("contact_phone","Optional","With country code if possible, e.g. +60 12 345 6789. Shown with Call and WhatsApp buttons."),
 ("contact_email","Optional","Shown with an Email button."),
 ("agency","Optional","Realtors only: your agency name. Leave blank for a private listing."),
 ("title","Optional","Short headline. Left blank, we write one, e.g. 3-bedroom house in Ipoh."),
 ("property_type","Optional","House, Condominium, Apartment, Terrace, Semi-detached, Bungalow, Villa, Townhouse, Land…"),
 ("bedrooms","Optional","Whole number. Studio = 0."),
 ("bathrooms","Optional","Whole number."),
 ("built_up","Optional","Floor area as a number. Units can go in the cell (1,200 sqft) or in area_unit."),
 ("land_size","Optional","Land or plot area as a number."),
 ("area_unit","Optional","sq ft, sq m, acres or hectares."),
 ("tenure","Optional","Freehold, Leasehold, or your local term (e.g. 99-year leasehold, Strata)."),
 ("deposit","Optional","Rentals: the deposit amount."),
 ("furnished","Optional","Rentals: Furnished, Partly furnished or Unfurnished."),
 ("available_from","Optional","Rentals: a date (2026-12-01) or Now."),
 ("minimum_lease_months","Optional","Rentals: 6 or more."),
 ("description","Optional","Up to 2,000 characters."),
 ("image_urls","Optional","Up to 12 photo links (https://…), separated by spaces or |. Or use columns image_1, image_2…"),
 ("reference","Optional","Your own listing ID, so you can match listings back to your system."),
 ("listing_url","Optional","Link to the listing on your own website."),
]
REALTOR_EXAMPLES = [
 {"price":"450000","currency":"MYR","location":"Ipoh, Perak, Malaysia","latitude":"4.5975","longitude":"101.0901","contact_name":"Aisha Rahman","contact_phone":"+60 12 345 6789","contact_email":"aisha@example.com","property_type":"Terrace","bedrooms":"3","bathrooms":"2","built_up":"1650","area_unit":"sq ft","tenure":"Freehold","description":"Corner unit, renovated kitchen, near schools.","image_urls":"https://example.com/photo1.jpg | https://example.com/photo2.jpg"},
 {"listing_type":"rent","price":"3800","currency":"SGD","latitude":"1.3521","longitude":"103.8198","contact_name":"Daniel Tan","contact_phone":"+65 9123 4567","agency":"Example Realty Pte Ltd","title":"2-bedroom condo with pool view","property_type":"Condominium","bedrooms":"2","bathrooms":"2","built_up":"850","area_unit":"sq ft","deposit":"7600","furnished":"Furnished","available_from":"2026-12-01","minimum_lease_months":"12","reference":"ER-1042","listing_url":"https://example.com/listings/1042"},
 {"price":"£325,000","location":"Bath, Somerset, United Kingdom","contact_name":"Sarah Jones","contact_email":"sarah@example.com","property_type":"Cottage","bedrooms":"2","bathrooms":"1","land_size":"0.25","area_unit":"acres","tenure":"Freehold","description":"Stone cottage with garden."},
]

SELLER = [
 ("For sale or rent","Optional","Sale or Rent. Blank means sale. " + LONG_TERM),
 ("Asking price","Required","Just the number, e.g. 450000 (or the monthly rent, e.g. 2500). RM 450,000 also works."),
 ("Location","Required","Town and country, e.g. Ipoh, Perak, Malaysia. Or leave blank and fill Latitude and Longitude."),
 ("Your name","Required","Shown to buyers and tenants so they know who they're talking to."),
 ("Phone or WhatsApp","Optional","With country code if you can, e.g. +60 12 345 6789."),
 ("Email","Optional","Buyers and tenants can email you."),
 ("Currency","Optional","MYR, SGD, AUD, GBP… Worked out from the location if left blank."),
 ("Property type","Optional","House, Condominium, Apartment, Terrace, Bungalow, Land…"),
 ("Bedrooms","Optional","A number. Studio = 0."),
 ("Bathrooms","Optional","A number."),
 ("Floor area","Optional","A number, e.g. 1200. Put the unit in Area unit."),
 ("Land size","Optional","A number, for houses and land."),
 ("Area unit","Optional","sq ft, sq m, acres or hectares."),
 ("Tenure","Optional","Freehold, Leasehold, or your local term."),
 ("Deposit","Optional","Rentals only: the deposit amount."),
 ("Furnished","Optional","Rentals only: Furnished, Partly furnished or Unfurnished."),
 ("Available from","Optional","Rentals only: a date, or Now."),
 ("Minimum lease (months)","Optional","Rentals only: 6 or more."),
 ("Headline","Optional","A short title. Leave blank and we write one, e.g. 3-bedroom house in Ipoh."),
 ("Description","Optional","Anything buyers or tenants should know."),
 ("Photo links","Optional","Links to photos already online (https://…), separated by spaces. Or add photos from your phone with the one-page form instead."),
 ("Latitude","Optional","Map pin, e.g. 4.5975. Tip: in Google Maps, press and hold on your home and copy the numbers."),
 ("Longitude","Optional","Map pin, e.g. 101.0901."),
]
SELLER_EXAMPLE = {"Asking price":"450000","Location":"Ipoh, Perak, Malaysia","Your name":"Aisha Rahman","Phone or WhatsApp":"+60 12 345 6789","Property type":"Terrace","Bedrooms":"3","Bathrooms":"2","Floor area":"1650","Area unit":"sq ft","Tenure":"Freehold","Description":"Corner unit, renovated kitchen, near schools.","Latitude":"4.5975","Longitude":"101.0901"}

def build(cols, examples, xlsx, csv_path, title, required_fill, lists):
    headers = [c[0] for c in cols]
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerow(headers)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Listings"
    for i, (h, r, d) in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=required_fill if r.startswith("Required") else "3F3F46")
        c.comment = Comment(f"{r}. {d}", "sellmyhouse.asia")
        ws.column_dimensions[c.column_letter].width = max(14, len(h) + 4)
    ws.freeze_panes = "A2"
    for header, options in lists.items():
        col = headers.index(header) + 1
        letter = ws.cell(1, col).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{letter}2:{letter}1000")
    g = wb.create_sheet("Guide")
    g["A1"] = f"{title} · sellmyhouse.asia"; g["A1"].font = Font(bold=True, size=14)
    g["A2"] = "One property per row on the Listings tab (only that tab is read). Save, then upload at sellmyhouse.asia/upload. Free, no account."
    g["A3"] = LONG_TERM
    g.append([]); g.append(["Column", "Needed?", "What to put in it"])
    for c in g[5]: c.font = Font(bold=True)
    for row in cols: g.append(list(row))
    g.column_dimensions["A"].width = 24; g.column_dimensions["B"].width = 12; g.column_dimensions["C"].width = 100
    e = wb.create_sheet("Example"); e.append(headers)
    for c in e[1]: c.font = Font(bold=True)
    for ex in examples: e.append([ex.get(h, "") for h in headers])
    wb.properties.title = title
    wb.save(xlsx)
    return headers

realtor_headers = build(REALTOR, REALTOR_EXAMPLES, "public/listing-template.xlsx", "public/listing-template.csv",
    "Global Database Format for Realtors", "C2410C",
    {"listing_type": ["sale", "rent"], "area_unit": ["sq ft", "sq m", "acres", "hectares"], "furnished": ["Furnished", "Partly furnished", "Unfurnished"]})
with open("public/listing-example.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(realtor_headers)
    for ex in REALTOR_EXAMPLES: w.writerow([ex.get(h, "") for h in realtor_headers])
build(SELLER, [SELLER_EXAMPLE], "public/seller-template.xlsx", "public/seller-template.csv",
    "Seller's Format for Individual Sellers", "15803D",
    {"For sale or rent": ["Sale", "Rent"], "Area unit": ["sq ft", "sq m", "acres", "hectares"], "Furnished": ["Furnished", "Partly furnished", "Unfurnished"]})

def ts(name, cols):
    out = [f"export const {name}: FormatColumn[] = ["]
    for n, r, d in cols:
        out.append(f"  {{ name: {json.dumps(n)}, required: {'true' if r.startswith('Required') else 'false'}, help: {json.dumps(d, ensure_ascii=False)} }},")
    out.append("];")
    return "\n".join(out)

with open("src/lib/listingFormat.ts", "w") as f:
    f.write("""// The published upload formats. Generated by scripts/make-templates.py — edit
// that script, not this file. Column names are matched loosely by
// uploadParse.ts, so these are the recommended names, not the only ones accepted.
export type FormatColumn = { name: string; required: boolean; help: string };

// Global Database Format for Realtors (public/listing-template.xlsx and .csv).
""" + ts("LISTING_COLUMNS", REALTOR) + """

// Seller's Format for Individual Sellers (public/seller-template.xlsx and .csv).
""" + ts("SELLER_COLUMNS", SELLER) + "\n")
print("ok")
